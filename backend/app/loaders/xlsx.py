"""Load form definitions from the challenge workbook or a CSV with the same columns.

Column names follow sheet `Data_Minimiser` of `data/klotenhack_challenge_datasets_hackers_v1_1.xlsx`:
form_id, form_name, business_context, field_order, field_name, field_label, field_type, required,
data_category, sensitive_flag, purpose_stated, purpose_text, retention_days, system_destination,
third_party_shared, risk_hint, jury_expected_flag, jury_expected_action, jury_expected_reason, inclusivity_note.

Jury columns are optional and are never given to the checks or the AI; they only feed `scripts/eval.py`.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Optional

from app.models import FieldSpec, FieldType, FormSchema

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_WORKBOOK = REPO_ROOT / "data" / "klotenhack_challenge_datasets_hackers_v1_1.xlsx"
SHEET = "Data_Minimiser"

_ENTRY_KEYWORDS = ("signup", "sign-up", "sign up", "onboarding", "demo request", "application",
                   "intake", "registration", "lead", "request")
_TYPES: set[str] = {"text", "textarea", "email", "number", "date", "dropdown", "file", "oauth", "checkbox"}


@dataclass
class JuryLabel:
    form_id: str
    field_id: str
    flag: bool
    action: str
    reason: str


def _yes(v: Any) -> Optional[bool]:
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in ("y", "yes", "true", "1"):
        return True
    if s in ("n", "no", "false", "0"):
        return False
    return None


def _int(v: Any) -> Optional[int]:
    if v is None or str(v).strip() == "":
        return None
    try:
        return int(float(v))
    except ValueError:
        return None


def _str(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _type(v: Any) -> FieldType:
    s = (str(v or "").strip().lower())
    return s if s in _TYPES else "other"  # type: ignore[return-value]


def derive_stage(name: str, business_context: str = "") -> str:
    hay = f"{name} {business_context}".lower()
    return "entry" if any(k in hay for k in _ENTRY_KEYWORDS) else "later"


def field_from_row(row: dict[str, Any]) -> FieldSpec:
    purpose_stated = _yes(row.get("purpose_stated"))
    purpose_text = _str(row.get("purpose_text"))
    if purpose_stated is False:
        purpose_text = None
    name = _str(row.get("field_name")) or _str(row.get("field_label")) or "field"
    return FieldSpec(
        field_id=name,
        order=_int(row.get("field_order")) or 0,
        name=name,
        label=_str(row.get("field_label")) or name,
        type=_type(row.get("field_type")),
        required=bool(_yes(row.get("required"))),
        data_category=_str(row.get("data_category")),
        sensitive=_yes(row.get("sensitive_flag")),
        purpose_text=purpose_text,
        retention_days=_int(row.get("retention_days")),
        destination=_str(row.get("system_destination")),
        third_party_shared=_yes(row.get("third_party_shared")),
        inclusivity_note=_str(row.get("inclusivity_note")),
    )


def forms_from_rows(rows: Iterable[dict[str, Any]], *, source: str = "demo") -> tuple[list[FormSchema], list[JuryLabel]]:
    forms: dict[str, FormSchema] = {}
    jury: list[JuryLabel] = []
    for row in rows:
        if not any(v not in (None, "") for v in row.values()):
            continue
        form_id = _str(row.get("form_id")) or "F000"
        if form_id not in forms:
            name = _str(row.get("form_name")) or form_id
            ctx = _str(row.get("business_context")) or ""
            forms[form_id] = FormSchema(
                form_id=form_id, name=name, business_context=ctx,
                stage=derive_stage(name, ctx), source=source,  # type: ignore[arg-type]
            )
        field = field_from_row(row)
        forms[form_id].fields.append(field)
        action = _str(row.get("jury_expected_action"))
        if action:
            jury.append(JuryLabel(
                form_id=form_id, field_id=field.field_id,
                flag=bool(_yes(row.get("jury_expected_flag"))),
                action=action, reason=_str(row.get("jury_expected_reason")) or "",
            ))
    for f in forms.values():
        f.fields.sort(key=lambda x: x.order)
    return list(forms.values()), jury


def rows_from_workbook(path: Path = DEFAULT_WORKBOOK, sheet: str = SHEET) -> list[dict[str, Any]]:
    import openpyxl  # local import: keeps CSV path free of the dependency

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    return [dict(zip(header, r)) for r in it]


def rows_from_csv(text: str) -> list[dict[str, Any]]:
    return list(csv.DictReader(io.StringIO(text)))


def load_workbook(path: Path = DEFAULT_WORKBOOK) -> tuple[list[FormSchema], list[JuryLabel]]:
    return forms_from_rows(rows_from_workbook(path), source="demo")


def load_demo_forms(path: Path = DEFAULT_WORKBOOK) -> list[FormSchema]:
    return load_workbook(path)[0]


def load_upload(filename: str, content: bytes) -> list[FormSchema]:
    """Parse an uploaded CSV or XLSX in the challenge column format."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=True) as tmp:
            tmp.write(content)
            tmp.flush()
            rows = rows_from_workbook(Path(tmp.name))
    else:
        rows = rows_from_csv(content.decode("utf-8-sig"))
    forms, _ = forms_from_rows(rows, source="upload")
    return forms
